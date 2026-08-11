from docx import Document

job_title = "Platform Engineer"

doc = Document()
doc.add_heading(f"Tailored Resume - {job_title}", level=1)
doc.add_paragraph("Linux, Docker, Jenkins, AWS, Kubernetes, Terraform, CI/CD")
doc.add_paragraph("Open to remote opportunities worldwide and willing to relocate for visa sponsorship.")
doc.save("Tailored_Resume.docx")

print("Generated Tailored_Resume.docx")
