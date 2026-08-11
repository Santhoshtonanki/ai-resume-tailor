import argparse
from pathlib import Path
from docx import Document

RESUME_MAP = {
    "devops": "resumes/Santhosh_DevOps_Resume.docx",
    "cloud": "resumes/Santhosh_Cloud_Resume.docx",
    "linux": "resumes/Santhosh_LinuxAdmin_Resume.docx",
    "noc": "resumes/Santhosh_NOC_Resume.docx",
}

KEYWORDS = {
    "devops": ["kubernetes", "docker", "jenkins", "terraform", "ci/cd", "aws"],
    "cloud": ["aws", "cloud", "ec2", "s3", "iam", "support"],
    "linux": ["linux", "shell", "bash", "systemd", "nginx"],
    "noc": ["monitoring", "ticket", "incident", "noc", "support"],
}

def choose_resume(text: str):
    text = text.lower()
    scores = {}
    for role, words in KEYWORDS.items():
        scores[role] = sum(1 for w in words if w in text)
    role = max(scores, key=scores.get)
    return role, RESUME_MAP[role], scores

def generate(role, source_resume, jd_text):
    doc = Document(source_resume)
    doc.add_page_break()
    doc.add_heading("Tailoring Notes", level=1)
    doc.add_paragraph(f"Selected role: {role.upper()}")
    doc.add_paragraph("Matched keywords from the job description:")
    found = []
    for words in KEYWORDS.values():
        for w in words:
            if w.lower() in jd_text.lower() and w not in found:
                found.append(w)
    for w in found:
        doc.add_paragraph(f"• {w}")
    out = Path("generated/Tailored_Resume.docx")
    doc.save(out)
    print(f"Saved {out}")

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", help="Path to JD text file")
    args = parser.parse_args()

    if args.file:
        jd = Path(args.file).read_text(encoding="utf-8")
    else:
        print("Paste job description, then press Enter twice:")
        lines = []
        while True:
            line = input()
            if line == "":
                break
            lines.append(line)
        jd = "\n".join(lines)

    role, resume, scores = choose_resume(jd)
    print(f"Selected: {role} -> {resume}")
    print(f"Scores: {scores}")
    generate(role, resume, jd)

if __name__ == "__main__":
    main()
